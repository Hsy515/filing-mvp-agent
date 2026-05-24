"""
S4 - XLSX 模板无损注入
依赖: pip install openpyxl
"""

from __future__ import annotations

import logging
import re
from copy import copy
from pathlib import Path

logger = logging.getLogger(__name__)

PLACEHOLDER_RE = re.compile(r"\{\{(#?)(\w+)\}\}")
TABLE_ANCHOR_RE = re.compile(r"<<table:(\w+)>>")  # 例: <<table:supplier_list>>


def inject_xlsx(
    template_path: str | Path,
    output_path: str | Path,
    placeholder_mapping: dict[str, dict],
    strict_mode: bool = True,
) -> dict:
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise RuntimeError("请先安装 openpyxl: pip install openpyxl") from e

    template_path = Path(template_path)
    output_path = Path(output_path)

    if not template_path.exists():
        return {"output_path": str(output_path), "status": "failed",
                "warnings": [f"模板不存在: {template_path}"],
                "placeholders_filled": 0, "placeholders_missing": []}

    wb = load_workbook(template_path)
    filled = 0
    missing: list[str] = []
    warnings: list[str] = []

    # 1) 单值替换 + 表格锚点处理
    for ws in wb.worksheets:
        # 单值替换
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None or not isinstance(cell.value, str):
                    continue
                text = cell.value

                # 表格锚点
                anchor_match = TABLE_ANCHOR_RE.search(text)
                if anchor_match:
                    anchor_name = anchor_match.group(1)
                    # 在 placeholder_mapping 里找对应 type=table_rows 的项
                    table_mapping = _find_table_mapping(placeholder_mapping, anchor_name)
                    if table_mapping:
                        _expand_table_anchor(ws, cell, table_mapping)
                        filled += 1
                    continue

                # 单值占位符
                new_text = text
                for ph, mapping in _iter_valid_mappings(placeholder_mapping):
                    if mapping.get("type") == "table_rows":
                        continue
                    if ph in new_text:
                        from inject_docx import _render_value  # 复用渲染函数
                        new_text = new_text.replace(ph, _render_value(mapping))
                        filled += 1
                if new_text != text:
                    cell.value = new_text

    # 2) 残留检查
    residual = _scan_residual(wb)
    if residual:
        if strict_mode:
            return {"output_path": str(output_path), "status": "failed",
                    "warnings": [f"strict_mode 下未替换占位符: {residual}"],
                    "placeholders_filled": filled, "placeholders_missing": list(residual)}
        for ph in residual:
            warnings.append(f"占位符未提供映射,已留空: {ph}")
            _replace_in_all_cells(wb, ph, "")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    status = "partial" if missing or warnings else "success"
    return {
        "output_path": str(output_path),
        "status": status,
        "placeholders_filled": filled,
        "placeholders_missing": missing,
        "warnings": warnings,
    }


def _find_table_mapping(mapping: dict, anchor_name: str) -> dict | None:
    """在 mapping 里找 anchor=anchor_name 的 table_rows 项"""
    for ph, m in _iter_valid_mappings(mapping):
        if m.get("type") == "table_rows" and m.get("anchor") == anchor_name:
            return m
    return None


def _iter_valid_mappings(mapping: dict):
    for ph, item in mapping.items():
        if str(ph).startswith("$") or not isinstance(item, dict):
            continue
        yield ph, item


def _expand_table_anchor(ws, anchor_cell, mapping: dict) -> None:
    """
    把 <<table:xxx>> 所在行作为表头下一行(模板行)的位置,把 rows 数据依次写入。
    简化版: 假设锚点行的下一行就是模板行,模板行的列字段名以 {{field}} 形式出现。
    """
    rows = mapping.get("rows", [])
    if not rows:
        return

    anchor_row = anchor_cell.row
    anchor_col = anchor_cell.column
    # 清除锚点单元格
    anchor_cell.value = None

    # 模板行(下一行)
    template_row_idx = anchor_row + 1
    template_row_cells = []
    for cell in ws[template_row_idx]:
        if cell.value is None:
            template_row_cells.append((cell.column, None, None))
            continue
        m = PLACEHOLDER_RE.fullmatch(str(cell.value).strip())
        field_name = m.group(2) if m else None
        # 保留样式
        template_row_cells.append((cell.column, field_name, copy(cell._style) if hasattr(cell, "_style") else None))

    # 按数据行数依次写
    for i, row_data in enumerate(rows):
        target_row_idx = template_row_idx + i
        # 确保行存在
        if target_row_idx > ws.max_row:
            pass  # openpyxl 写入时自动扩展
        for col, field_name, style in template_row_cells:
            target_cell = ws.cell(row=target_row_idx, column=col)
            if field_name is not None and field_name in row_data:
                target_cell.value = row_data[field_name]
            if style is not None:
                target_cell._style = style


def _scan_residual(wb) -> set[str]:
    found = set()
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    for m in PLACEHOLDER_RE.finditer(cell.value):
                        found.add(m.group(0))
    return found


def _replace_in_all_cells(wb, placeholder: str, value: str) -> None:
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and placeholder in cell.value:
                    cell.value = cell.value.replace(placeholder, value)


if __name__ == "__main__":
    import argparse, json
    p = argparse.ArgumentParser()
    p.add_argument("--template", required=True)
    p.add_argument("--output", required=True)
    p.add_argument("--mapping", required=True)
    p.add_argument("--no-strict", action="store_true")
    args = p.parse_args()

    mapping = json.loads(Path(args.mapping).read_text(encoding="utf-8-sig"))
    res = inject_xlsx(args.template, args.output, mapping, strict_mode=not args.no_strict)
    print(json.dumps(res, ensure_ascii=False, indent=2))
