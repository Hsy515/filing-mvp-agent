"""
S4 - DOCX 模板无损注入
依赖: pip install python-docx
"""

from __future__ import annotations

import copy
import logging
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)

PLACEHOLDER_RE = re.compile(r"\{\{(#?)(\w+)\}\}")  # 例: {{project_name}} 或 {{#supplier_table}}


@dataclass
class InjectionResult:
    output_path: str
    status: str = "success"  # success | partial | failed
    placeholders_filled: int = 0
    placeholders_missing: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    audit: dict[str, Any] = field(default_factory=dict)

    def to_dict(self) -> dict:
        return {
            "output_path": self.output_path,
            "status": self.status,
            "placeholders_filled": self.placeholders_filled,
            "placeholders_missing": self.placeholders_missing,
            "warnings": self.warnings,
            "audit": self.audit,
        }


def inject_docx(
    template_path: str | Path,
    output_path: str | Path,
    placeholder_mapping: dict[str, dict],
    strict_mode: bool = True,
) -> InjectionResult:
    """主入口: 把 mapping 注入到 docx 模板,保存到 output_path。"""
    try:
        from docx import Document
    except ImportError as e:
        raise RuntimeError("请先安装 python-docx: pip install python-docx") from e

    template_path = Path(template_path)
    output_path = Path(output_path)
    result = InjectionResult(output_path=str(output_path))

    if not template_path.exists():
        result.status = "failed"
        result.warnings.append(f"模板不存在: {template_path}")
        return result

    doc = Document(template_path)

    # 1) 扫描所有占位符位置
    found = _scan_placeholders(doc)
    logger.info("扫描到占位符: %s", found)
    result.audit = _build_preflight_audit(found, placeholder_mapping)

    # 2) 单值占位符替换
    for placeholder, mapping in _iter_valid_mappings(placeholder_mapping, result.warnings):
        m = PLACEHOLDER_RE.fullmatch(placeholder)
        if not m:
            result.warnings.append(f"非法占位符格式: {placeholder}")
            continue
        is_table_row = m.group(1) == "#"
        if is_table_row:
            ok = _inject_table_rows(doc, placeholder, mapping)
        else:
            ok = _inject_single_value(doc, placeholder, mapping)

        if ok:
            result.placeholders_filled += 1
        else:
            logger.info("模板中未使用映射占位符: %s", placeholder)

    # 3) 残留占位符处理
    residual = _scan_placeholders(doc)
    if residual:
        result.placeholders_missing = sorted(residual)
        if strict_mode:
            result.status = "failed"
            result.warnings.append(f"strict_mode 下仍有未替换占位符: {residual}")
            return result
        # 非严格模式: 替换为空,但记录
        for ph in residual:
            _inject_single_value(doc, ph, {"type": "text", "value": ""})
            result.warnings.append(f"占位符未提供映射,已替换为空: {ph}")

    # 4) 保存
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)

    if result.placeholders_missing:
        result.status = "partial"
    return result


def audit_template(template_path: str | Path, placeholder_mapping: dict | None = None) -> dict:
    """只扫描模板,不写输出文件。用于模板设计师交付前自检。"""
    try:
        from docx import Document
    except ImportError as e:
        raise RuntimeError("请先安装 python-docx: pip install python-docx") from e

    doc = Document(template_path)
    found = _scan_placeholders(doc)
    return _build_preflight_audit(found, placeholder_mapping or {})


# ---------- 扫描 ----------

def _scan_placeholders(doc) -> set[str]:
    found = set()
    # 段落
    for p in doc.paragraphs:
        for m in PLACEHOLDER_RE.finditer(p.text):
            found.add(m.group(0))
    # 表格
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for m in PLACEHOLDER_RE.finditer(p.text):
                        found.add(m.group(0))
    # Header / Footer
    for section in doc.sections:
        for container in (section.header, section.footer):
            for p in container.paragraphs:
                for m in PLACEHOLDER_RE.finditer(p.text):
                    found.add(m.group(0))
    for text in _iter_xml_text_values(doc):
        for m in PLACEHOLDER_RE.finditer(text):
            found.add(m.group(0))
    return found


def _build_preflight_audit(found: set[str], placeholder_mapping: dict) -> dict:
    provided = {
        ph for ph, item in placeholder_mapping.items()
        if not str(ph).startswith("$") and isinstance(item, dict)
    }
    row_child_placeholders = _row_child_placeholders(placeholder_mapping)
    covered = provided | row_child_placeholders
    illegal = sorted(ph for ph in provided if not PLACEHOLDER_RE.fullmatch(ph))
    return {
        "placeholders_found": sorted(found),
        "mappings_provided": sorted(provided),
        "row_child_placeholders": sorted(row_child_placeholders),
        "missing_mappings": sorted(found - covered),
        "unused_mappings": sorted(provided - found),
        "illegal_mappings": illegal,
    }


def _row_child_placeholders(placeholder_mapping: dict) -> set[str]:
    child = set()
    for mapping in placeholder_mapping.values():
        if not isinstance(mapping, dict) or mapping.get("type") != "table_rows":
            continue
        for row in mapping.get("rows", []):
            if isinstance(row, dict):
                child.update(f"{{{{{key}}}}}" for key in row)
    return child


def _iter_valid_mappings(placeholder_mapping: dict, warnings: list[str]):
    for placeholder, mapping in placeholder_mapping.items():
        if str(placeholder).startswith("$"):
            continue
        if not isinstance(mapping, dict):
            warnings.append(f"忽略非占位符映射项: {placeholder}")
            continue
        yield placeholder, mapping


# ---------- 单值注入 ----------

def _inject_single_value(doc, placeholder: str, mapping: dict) -> bool:
    """把模板里所有出现的 placeholder 替换为 mapping 渲染结果"""
    value_str = _render_value(mapping)
    replaced = False

    for p in doc.paragraphs:
        if placeholder in p.text:
            _replace_in_paragraph(p, placeholder, value_str)
            replaced = True

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if placeholder in p.text:
                        _replace_in_paragraph(p, placeholder, value_str)
                        replaced = True

    for section in doc.sections:
        for container in (section.header, section.footer):
            for p in container.paragraphs:
                if placeholder in p.text:
                    _replace_in_paragraph(p, placeholder, value_str)
                    replaced = True

    if _replace_in_xml_text_nodes(doc, placeholder, value_str):
        replaced = True

    return replaced


def _replace_in_paragraph(paragraph, placeholder: str, value: str) -> None:
    """
    占位符可能横跨多个 Run。先合并到一个 Run,再做字符串替换。
    保留第一个 Run 的样式属性。
    """
    if placeholder not in paragraph.text:
        return

    runs = paragraph.runs
    if not runs:
        return

    # 找包含完整占位符的 run; 若都不完整,合并所有 run
    for run in runs:
        if placeholder in run.text:
            run.text = run.text.replace(placeholder, value)
            return

    # 跨 run: 合并全部到第一个,保留第一个 run 的样式
    full = "".join(r.text for r in runs)
    replaced = full.replace(placeholder, value)
    runs[0].text = replaced
    for r in runs[1:]:
        r.text = ""


def _iter_xml_text_values(doc):
    for node in _iter_xml_text_nodes(doc):
        if node.text:
            yield node.text


def _iter_xml_text_nodes(doc):
    from docx.oxml.ns import qn

    roots = [doc.element]
    for section in doc.sections:
        roots.append(section.header._element)
        roots.append(section.footer._element)

    seen = set()
    for root in roots:
        for node in root.iter(qn("w:t")):
            ident = id(node)
            if ident in seen:
                continue
            seen.add(ident)
            yield node


def _replace_in_xml_text_nodes(doc, placeholder: str, value: str) -> bool:
    """Handle placeholders stored in text boxes or drawing XML not exposed as paragraphs."""
    replaced = False
    for node in _iter_xml_text_nodes(doc):
        if node.text and placeholder in node.text:
            node.text = node.text.replace(placeholder, value)
            replaced = True
    return replaced


# ---------- 表格行注入 ----------

def _inject_table_rows(doc, placeholder: str, mapping: dict) -> bool:
    """
    定位 placeholder 所在的"模板行",复制该行 N 次,逐行填值。
    """
    rows_data = mapping.get("rows", [])
    if not rows_data:
        return False

    for table in doc.tables:
        template_row_idx = None
        for i, row in enumerate(table.rows):
            for cell in row.cells:
                if placeholder in cell.text:
                    template_row_idx = i
                    break
            if template_row_idx is not None:
                break

        if template_row_idx is None:
            continue

        template_row = table.rows[template_row_idx]
        new_rows_xml = []

        for row_data in rows_data:
            new_row = copy.deepcopy(template_row._tr)
            new_rows_xml.append((new_row, row_data))

        # 把新行依次插入到模板行之前
        for new_row_xml, row_data in new_rows_xml:
            template_row._tr.addprevious(new_row_xml)
            # 对新插入的行做单值替换 — 通过重新读 cells
            # 这里简化处理: 直接对 XML 字符串扫描子占位符 {{field_name}}
            from docx.oxml.ns import qn
            for child_p in new_row_xml.iter(qn("w:p")):
                # 拿到段落文本
                text_elements = list(child_p.iter(qn("w:t")))
                full_text = "".join((t.text or "") for t in text_elements)
                replaced = full_text
                replaced = replaced.replace(placeholder, str(row_data.get("index", "")))
                for k, v in row_data.items():
                    replaced = replaced.replace(f"{{{{{k}}}}}", str(v))
                if replaced != full_text and text_elements:
                    text_elements[0].text = replaced
                    for t in text_elements[1:]:
                        t.text = ""

        # 删除模板行
        template_row._tr.getparent().remove(template_row._tr)
        return True

    return False


# ---------- 值渲染 ----------

def _render_value(mapping: dict) -> str:
    vtype = mapping.get("type", "text")
    value = mapping.get("value", "")

    if vtype == "text":
        return str(value or "")
    if vtype == "number":
        try:
            num = float(value)
        except (TypeError, ValueError):
            return str(value)
        # 千分位 + 两位小数
        return f"{num:,.2f}"
    if vtype == "date":
        return _render_date(value, mapping.get("format", "YYYY-MM-DD"))
    return str(value or "")


def _render_date(value: str, fmt: str) -> str:
    from datetime import datetime
    if not value:
        return ""
    # 尝试解析
    for try_fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%Y/%m/%d"):
        try:
            dt = datetime.strptime(str(value).strip(), try_fmt)
            break
        except ValueError:
            continue
    else:
        return str(value)
    # 先处理时间,避免 HH:MM 里的 MM 被当作月份。
    out = fmt.replace("HH:MM", f"{dt.hour:02d}:{dt.minute:02d}")
    out = out.replace("HH:mm", f"{dt.hour:02d}:{dt.minute:02d}")
    out = out.replace("hh:mm", f"{dt.hour:02d}:{dt.minute:02d}")
    return (out.replace("YYYY", str(dt.year))
               .replace("DD", f"{dt.day:02d}")
               .replace("MM", f"{dt.month:02d}")
               .replace("D", str(dt.day))
               .replace("M", str(dt.month)))


if __name__ == "__main__":
    import argparse, json
    p = argparse.ArgumentParser()
    p.add_argument("--template", required=True)
    p.add_argument("--output")
    p.add_argument("--mapping", required=True, help="placeholder mapping JSON")
    p.add_argument("--no-strict", action="store_true")
    p.add_argument("--audit-only", action="store_true", help="只扫描模板占位符,不生成输出文件")
    args = p.parse_args()

    mapping = json.loads(Path(args.mapping).read_text(encoding="utf-8-sig"))
    if args.audit_only:
        print(json.dumps(audit_template(args.template, mapping), ensure_ascii=False, indent=2))
    else:
        if not args.output:
            p.error("--output is required unless --audit-only is set")
        res = inject_docx(args.template, args.output, mapping, strict_mode=not args.no_strict)
        print(json.dumps(res.to_dict(), ensure_ascii=False, indent=2))
